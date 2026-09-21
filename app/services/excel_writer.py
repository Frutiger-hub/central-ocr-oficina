from datetime import date,datetime
from pathlib import Path
from shutil import copy2
from openpyxl import load_workbook
class WorkbookError(Exception): pass
def _last_data_row(ws):
    for r in range(ws.max_row,1,-1):
        if ws.cell(r,1).value not in (None,""): return r
    return 1
def _existing_os(ws,os_number):
    target=str(os_number).strip()
    return any(ws.cell(r,1).value not in (None,"") and str(ws.cell(r,1).value).strip()==target for r in range(2,ws.max_row+1))
def _coerce_date(v): return date.fromisoformat(v) if v else None
def create_backup(path):
    d=path.parent/"backups"; d.mkdir(parents=True,exist_ok=True)
    p=d/f"{path.stem}_backup_{datetime.now():%Y%m%d_%H%M%S}{path.suffix}"; copy2(path,p); return p
def append_launch(workbook_path,os_number,today,client,technician,data_final,status,rebobinamento,color,power):
    if not workbook_path.exists(): raise WorkbookError(f"Planilha não encontrada: {workbook_path}")
    if not os_number or not client or not technician: raise WorkbookError("O.S., cliente e técnico são obrigatórios.")
    wb=load_workbook(workbook_path,keep_vba=True)
    try:
        for n in ("MONTADOS","ENROLADOR"):
            if n not in wb.sheetnames: raise WorkbookError(f"Aba {n} não encontrada.")
        m,e=wb["MONTADOS"],wb["ENROLADOR"]; backup=create_backup(workbook_path); written=[]; skipped=[]
        if _existing_os(m,os_number): skipped.append("MONTADOS")
        else:
            r=_last_data_row(m)+1
            vals=[os_number,today,client,technician,_coerce_date(data_final),status]
            for c,v in enumerate(vals,1): m.cell(r,c,v)
            m.cell(r,2).number_format="dd/mm/yyyy"
            if vals[4]: m.cell(r,5).number_format="dd/mm/yyyy"
            written.append("MONTADOS")
        if rebobinamento:
            if _existing_os(e,os_number): skipped.append("ENROLADOR")
            else:
                r=_last_data_row(e)+1
                vals=[os_number,today,client,technician,color or "ROSA",power or ""]
                for c,v in enumerate(vals,1): e.cell(r,c,v)
                e.cell(r,2).number_format="dd/mm/yyyy"; written.append("ENROLADOR")
        if written: wb.save(workbook_path)
        return {"written":written,"skipped":skipped,"backup":str(backup) if written else None}
    finally: wb.close()
